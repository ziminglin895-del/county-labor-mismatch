# -*- coding: utf-8 -*-
"""Join 2020 county-yearbook controls onto the census education table."""
import os
import shutil

import pandas as pd
import win32com.client
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

DESKTOP = r"D:\Desktop"
CENSUS = os.path.join(DESKTOP, "成员B_七普人口学历_真表.xlsx")
FILES = {
    "浙江": "2026_09_26_13_09_47_N2022040099000018.xls",
    "福建": "2026_09_26_13_10_10_N2022040099000022.xls",
    "江西": "2026_09_26_13_10_20_N2022040099000023.xls",
    "河南": "2026_09_26_13_10_28_N2022040099000027.xls",
    "四川": "2026_09_26_13_10_36_N2022040099000040.xls",
    "安徽福建交界": "2026_09_26_13_31_47_N2022040099000021.xls",
    "江西山东交界": "2026_09_26_13_31_25_N2022040099000024.xls",
    "河南湖北交界": "2026_09_26_13_31_40_N2022040099000028.xls",
    "四川贵州交界": "2026_09_26_13_32_14_N2022040099000041.xls",
    "浙江安徽交界": "2026_09_26_15_05_42_N2022040099000019.xls",
    "山东河南交界": "2026_09_26_15_05_26_N2022040099000026.xls",
}
INDICATORS = {
    "户籍人口": "hukou_pop_2020",
    "地区生产总值": "gdp_wanyuan_2020",
    "第一产业增加值": "gdp_primary_wanyuan_2020",
    "第二产业增加值": "gdp_secondary_wanyuan_2020",
    "第三产业增加值": "gdp_tertiary_wanyuan_2020",
    "地方一般公共预算收入": "fiscal_rev_wanyuan_2020",
    "地方一般公共预算支出": "fiscal_exp_wanyuan_2020",
    "规模以上工业企业": "scale_industrial_firms_2020",
}
OUT_PROJ = r"D:\资料\经管学业\县域劳动力错配研究\成员B_宏观学历\data\成员B_学历与宏观_合并交付.xlsx"
OUT_DESK = os.path.join(DESKTOP, "成员B_学历与宏观_合并交付.xlsx")


def norm_name(value):
    if value is None:
        return ""
    text = str(value).replace(" ", "").replace("\u3000", "").replace("\n", "").replace("\r", "")
    return text.strip()


def to_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("，", "")
    if text in {"", "-", "—", "…", ".", "·"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def extract_file(excel, path, province):
    wb = excel.Workbooks.Open(path)
    ws = wb.Sheets(1)
    used = ws.UsedRange.Value
    wb.Close(False)
    rows = []
    current = None
    for raw in used:
        cells = list(raw)
        label = norm_name(cells[0] if len(cells) else "")
        if label in {"指标", "指 标"}:
            names = [norm_name(x) for x in cells[2:]]
            names = [n for n in names if n]
            current = names
            continue
        if current is None or label not in INDICATORS:
            continue
        values = cells[2:2 + len(current)]
        for name, value in zip(current, values):
            rows.append((province, name, INDICATORS[label], to_number(value)))
    return rows


def main():
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    long_rows = []
    try:
        for province, fn in FILES.items():
            got = extract_file(excel, os.path.join(DESKTOP, fn), province)
            print(province, "cells", len(got), "counties", len({r[1] for r in got}))
            long_rows.extend(got)
    finally:
        excel.Quit()

    long_df = pd.DataFrame(long_rows, columns=["province_file", "yb_name", "variable", "value"])
    wide = long_df.pivot_table(
        index=["province_file", "yb_name"], columns="variable", values="value", aggfunc="first"
    ).reset_index()

    census = pd.read_excel(CENSUS, sheet_name="五省412县", dtype={"county_code": str, "county_code_2020census": str})
    census["county_code"] = census["county_code"].astype(str).str.replace(r"\.0$", "", regex=True).str.zfill(6)
    census["county_code_2020census"] = census["county_code_2020census"].astype(str).str.replace(r"\.0$", "", regex=True).str.zfill(6)
    for col in ["county_name", "name_2020census"]:
        census[col + "_key"] = census[col].map(norm_name)

    by_now = census.drop_duplicates("county_name_key").set_index("county_name_key")["county_code"]
    by_old = census.drop_duplicates("name_2020census_key").set_index("name_2020census_key")["county_code"]

    def match_code(name):
        if name in by_old.index:
            return by_old.loc[name]
        if name in by_now.index:
            return by_now.loc[name]
        return None

    wide["county_code"] = wide["yb_name"].map(match_code)
    matched = wide[wide["county_code"].notna()].copy()
    dupes = matched[matched["county_code"].duplicated(keep=False)]
    if len(dupes):
        print("DUPLICATE CODES")
        print(dupes[["province_file", "yb_name", "county_code"]].to_string(index=False))
    matched = matched.drop_duplicates("county_code", keep="first")

    sample_codes = set(census["county_code"])
    unmatched_yb = wide[wide["county_code"].isna()].copy()
    # districts outside the sample are expected; counties and county-level cities are not
    unmatched_yb["kind"] = unmatched_yb["yb_name"].map(
        lambda n: "区" if n.endswith("区") else ("市" if n.endswith("市") else ("县" if n.endswith("县") or n.endswith("旗") else "其他"))
    )
    missing_sample = census[~census["county_code"].isin(set(matched["county_code"]))].copy()

    keep_cols = ["county_code"] + list(INDICATORS.values())
    merged = census.drop(columns=["county_name_key", "name_2020census_key"]).merge(
        matched[keep_cols], on="county_code", how="left"
    )
    merged["share_secondary_gdp"] = merged["gdp_secondary_wanyuan_2020"] / merged["gdp_wanyuan_2020"]
    merged["share_tertiary_gdp"] = merged["gdp_tertiary_wanyuan_2020"] / merged["gdp_wanyuan_2020"]
    merged["gdp_per_capita_yuan"] = merged["gdp_wanyuan_2020"] * 10000 / merged["pop_total_2020"]
    for col in ["share_secondary_gdp", "share_tertiary_gdp", "gdp_per_capita_yuan", "share_college_census"]:
        merged[col] = merged[col].round(6)
    merged["edu_years_gap"] = merged["edu_years_gap"].round(2)

    print("sample", len(merged), "yearbook matched", int(merged["gdp_wanyuan_2020"].notna().sum()))
    print("missing gdp", merged.loc[merged["gdp_wanyuan_2020"].isna(), ["province", "county_name"]].to_string(index=False))
    print("unmatched county-like", unmatched_yb.loc[unmatched_yb["kind"] != "区", ["province_file", "yb_name", "kind"]].to_string(index=False))
    print("districts left out", int((unmatched_yb["kind"] == "区").sum()))

    spots = merged[merged["county_name"].isin(["固始县", "布拖县", "义乌市", "淳安县", "沙县区", "会理市", "龙南市", "偃师区"])]
    print(spots[["county_name", "county_code", "pop_total_2020", "edu_years_census", "gdp_wanyuan_2020", "gdp_per_capita_yuan", "scale_industrial_firms_2020"]].to_string(index=False))

    note = pd.DataFrame({
        "项": [
            "能不能交差",
            "学历能不能当求职者",
            "年鉴来源",
            "交叉页",
            "接上了什么",
            "规上缺6县",
            "年鉴没有的",
            "人口用哪个人口",
            "人均GDP",
            "产业占比",
            "市辖区",
            "区划",
            "缺失",
        ],
        "说明": [
            "五省 412 个县的行都在。学历、常住人口、GDP、三次产业增加值、一般公共预算收支 412 县都有数。规上工业企业缺 6 个县，见下一条。二三产业从业人数这本县市卷没有这一行，空着。",
            "edu_years_census 和 share_college_census 是2020年6岁及以上常住人口，在校生算在正在读的那一档。只能当县域人力资本存量，不能写成2026年5月的求职者学历。",
            "《中国县域统计年鉴2021》县市卷，数据年份2020。知网五省整省 Excel，加上浙皖、皖闽、赣鲁、鲁豫、豫鄂、川黔六个交界页。",
            "整省文件不含印在交界页上的县。浙皖、皖闽、赣鲁、鲁豫、豫鄂、川黔六个交界页已并入，412 个样本县全部对上。",
            "地区生产总值、一二三产业增加值、一般公共预算收入和支出，单位都是万元。规模以上工业企业单位是个。户籍人口是户籍，不是常住人口。",
            "scale_industrial_firms_2020 有 6 个县为空：道孚县、新龙县、德格县、石渠县、色达县、稻城县，都在甘孜藏族自治州。年鉴该格未刊载，不是 0。乡城县规上是 7，不是缺失。回归时这 6 个县作缺失处理，剔除或单独标记，不要填 0。",
            "没有第二产业从业人员，也没有第三产业从业人员。空着，没有填0，没有用长表抽样或随机数代替。",
            "pop_total_2020 是七普常住人口。年鉴里的户籍人口单列，不替换常住人口。",
            "gdp_per_capita_yuan = 地区生产总值（万元）×10000 / 七普常住人口。分母不是户籍人口。",
            "share_secondary_gdp、share_tertiary_gdp 是增加值除以地区生产总值，不是从业人数占比。",
            "年鉴里的市辖区没有并进样本。撤县设区的沙县区、偃师区、孟津区按2020年旧名接上。",
            "county_code 用现行码。龙南在年鉴和普查里都已是龙南市360783。会理年鉴名是会理县，现行码513402。金门没有普查，不在表里。",
            "样本 412 县的 GDP 没有空值。『年鉴没对上的县』应是空表。交界页上其他省的县没有并进来。",
        ],
    })
    miss = missing_sample[["province", "county_code", "county_name", "name_2020census"]].sort_values(["province", "county_code"])
    trial = merged[merged["province"].isin(["河南", "浙江", "四川"])].copy()

    def write(path, sheets):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            for name, df in sheets:
                df.to_excel(writer, sheet_name=name[:31], index=False)
        wb = load_workbook(path)
        fill = PatternFill("solid", fgColor="1F4E79")
        font = Font(color="FFFFFF", bold=True, name="Calibri")
        for ws in wb.worksheets:
            for cell in ws[1]:
                cell.fill = fill
                cell.font = font
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for col in ws.columns:
                header = str(col[0].value or "")
                if "code" in header:
                    for cell in col[1:]:
                        if cell.value not in (None, ""):
                            cell.number_format = "@"
                            cell.value = str(cell.value).replace(".0", "").zfill(6)
                ws.column_dimensions[col[0].column_letter].width = min(max(len(header), 14), 42)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.row_dimensions[1].height = 22
        wb.save(path)

    sheets = [
        ("试分析三省286县", trial),
        ("五省412县", merged),
        ("年鉴没对上的县", miss),
        ("读表说明", note),
    ]
    write(OUT_PROJ, sheets)
    shutil.copy2(OUT_PROJ, OUT_DESK)
    print("wrote", OUT_DESK)
    print("trial", len(trial), "full", len(merged), "miss", len(miss))


if __name__ == "__main__":
    main()
